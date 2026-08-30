import os
import pandas as pd
from sklearn import set_config
from sklearn.linear_model import LogisticRegression
from sklearn.naive_bayes import GaussianNB
from sklearn.tree import DecisionTreeClassifier
from sklearn.svm import SVC
from sklearn.calibration import CalibratedClassifierCV
from sklearn.discriminant_analysis import LinearDiscriminantAnalysis
from sklearn.ensemble import RandomForestClassifier
from sklearn.neighbors import KNeighborsClassifier
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import LeaveOneOut, cross_val_predict
from sklearn.metrics import (precision_recall_curve, roc_curve, roc_auc_score, average_precision_score,
    confusion_matrix, precision_score,)
from lightgbm import LGBMClassifier
from xgboost import XGBClassifier
from catboost import CatBoostClassifier
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl.styles import PatternFill

# Classic-ml LOO-CV on the matched-pairs dataset built (142 rows: 71 core "case" + 71 confound-matched "control"). 
# Three feature-set variants isolate what loop-boundary info adds: noloop (grantham + third_nt/minus1_nt context only) -> loopflag
# (+ binary in_loop_boundary, a case/control separator by matching construction) -> withloop (+ continuous length/stem_length/pin_energy).

set_config(transform_output="pandas")

IN_TSV = "matched_pairs_dataset.tsv"
RESULTS_DIR = "results_classic_ml_matched_pairs"
os.makedirs(RESULTS_DIR, exist_ok=True)

def out_path(filename):
    return os.path.join(RESULTS_DIR, filename)

matched = pd.read_csv(IN_TSV, sep='\t')
n_case, n_ctrl = (matched.group == "case").sum(), (matched.group == "control").sum()
print(f"loaded {IN_TSV}: {len(matched)} rows ({n_case} case / {n_ctrl} control)")

THIRD_NT_COLS = ["third_nt_A", "third_nt_C", "third_nt_G", "third_nt_T"]
MINUS1_NT_COLS = ["minus1_nt_A", "minus1_nt_C", "minus1_nt_G", "minus1_nt_T"]
LOOP_BOUNDARY_COLS = ["loop_boundary_length", "loop_boundary_stem_length", "loop_boundary_pin_energy"]

# ds1_mutation excluded: only 1 positive among the 71 core rows
TARGETS = ["any_mutation", "ds2_mutation", "ds3_mutation"]

FEATURES_NO_LOOP = ["grantham"] + THIRD_NT_COLS + MINUS1_NT_COLS
FEATURES_LOOP_FLAG_ONLY = FEATURES_NO_LOOP + ["in_loop_boundary"]
FEATURES_WITH_LOOP = FEATURES_NO_LOOP + ["in_loop_boundary"] + LOOP_BOUNDARY_COLS
FEATURE_SETS = {"noloop": FEATURES_NO_LOOP, "loopflag": FEATURES_LOOP_FLAG_ONLY, "withloop": FEATURES_WITH_LOOP}


def make_models(y_true):
    n_pos = int(y_true.sum())
    n_neg = len(y_true) - n_pos
    scale_pos_weight = n_neg / n_pos

    return {
        "LogisticRegression": LogisticRegression(max_iter=1000, class_weight="balanced"),
        "NaiveBayes": GaussianNB(),
        #"DecisionTree": DecisionTreeClassifier(random_state=1, class_weight="balanced"),
        "SVM": CalibratedClassifierCV(SVC(kernel="rbf", random_state=1, class_weight="balanced"), ensemble=False),
        "LDA": LinearDiscriminantAnalysis(),
        "RandomForest": RandomForestClassifier(random_state=1, class_weight="balanced"),
        #"KNeighbors": KNeighborsClassifier(),
        "LightGBM": LGBMClassifier(random_state=1, class_weight="balanced",
                                    n_estimators=100, max_depth=3, verbose=-1, n_jobs=1),
        "XGBoost": XGBClassifier(random_state=1, scale_pos_weight=scale_pos_weight,
                                  n_estimators=100, max_depth=3, eval_metric="logloss",
                                  verbosity=0, n_jobs=1),
        "CatBoost": CatBoostClassifier(random_state=1, auto_class_weights="Balanced",
                                        iterations=100, depth=3, verbose=False,
                                        allow_writing_files=False, thread_count=1),
    }


def save_predictions_xlsx(out, target, path):
    def err_type(r):
        if r["pred_class"] == 1 and r[target] == 0: return "FP"
        if r["pred_class"] == 0 and r[target] == 1: return "FN"
        return ""
    out = out.copy()
    out["error"] = out.apply(err_type, axis=1)
    out.to_excel(path, index=False)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    fp_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    err_col = out.columns.get_loc("error") + 1
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row=row, column=err_col).value
        fill = fp_fill if val == "FP" else fn_fill if val == "FN" else None
        if fill:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill
    wb.save(path)
    return path


def run_for_feature_set(tag, feature_cols, target):
    X = matched[feature_cols]
    y_true = matched[target].values
    run_tag = f"{tag}_{target}"

    summary_rows = []
    roc_lines = []  # (line, roc_auc) -- legend sorted by roc_auc descending after the loop
    pr_lines = []   # (line, pr_auc)
    fig_roc, ax_roc = plt.subplots(figsize=(6, 5))
    fig_pr, ax_pr = plt.subplots(figsize=(6, 5))

    for model_name, estimator in make_models(y_true).items():
        pipe = Pipeline([
            ("scaler", StandardScaler()),
            ("clf", estimator),
        ])

        prob = cross_val_predict(pipe, X, y_true, cv=LeaveOneOut(), method="predict_proba", n_jobs=-1)[:, 1]

        threshold = prob[y_true == 1].min()
        y_pred = (prob >= threshold).astype(int)

        tn, fp, fn, tp = confusion_matrix(y_true, y_pred).ravel()
        recall = tp / (tp + fn)
        precision = precision_score(y_true, y_pred, zero_division=0)
        roc_auc = roc_auc_score(y_true, prob)
        pr_auc = average_precision_score(y_true, prob)

        odf = matched.copy()
        odf["pred_class"] = y_pred
        odf["pred_prob"] = prob

        odf.to_csv(out_path(f"predictions_{model_name}_{run_tag}.csv"), index=False)
        save_predictions_xlsx(odf, target, out_path(f"predictions_{model_name}_{run_tag}.xlsx"))

        fpr, tpr, _ = roc_curve(y_true, prob)
        (roc_line,) = ax_roc.plot(fpr, tpr, label=f"{model_name} (AUC={roc_auc:.3f})")
        roc_lines.append((roc_line, roc_auc))

        prec, rec, _ = precision_recall_curve(y_true, prob)
        (pr_line,) = ax_pr.plot(rec, prec, label=f"{model_name} (AP={pr_auc:.3f})")
        pr_lines.append((pr_line, pr_auc))
        ax_pr.scatter([recall], [precision], zorder=5)

        summary_rows.append({
            "target": target, "feature_set": tag, "n_rows": len(matched), "n_features": len(feature_cols),
            "n_positives": int(y_true.sum()),
            "model": model_name, "roc_auc": roc_auc, "pr_auc": pr_auc,
            "threshold": threshold, "recall": recall, "precision": precision,
            "TP": tp, "FP": fp, "FN": fn, "TN": tn,
        })

    ax_roc.plot([0, 1], [0, 1], "--", color="grey")
    ax_roc.set_xlabel("False positive rate"); ax_roc.set_ylabel("True positive rate")
    #ax_roc.set_title(f"ROC — classic ML models, matched pairs, {run_tag} (leave-one-out OOF)")
    roc_lines.sort(key=lambda t: t[1], reverse=True)
    ax_roc.legend([l for l, _ in roc_lines], [l.get_label() for l, _ in roc_lines], loc="lower right", fontsize=8)
    fig_roc.savefig(out_path(f"roc_curve_all_models_{run_tag}.png"), dpi=600, bbox_inches="tight")
    plt.close(fig_roc)

    pr_lines.sort(key=lambda t: t[1], reverse=True)
    ax_pr.set_xlabel("Recall"); ax_pr.set_ylabel("Precision")
    ax_pr.set_title(f"PR — classic ML models, matched pairs, {run_tag} (dots = working point at recall=1)")
    ax_pr.legend([l for l, _ in pr_lines], [l.get_label() for l, _ in pr_lines], loc="lower left", fontsize=8)
    fig_pr.savefig(out_path(f"pr_curve_all_models_{run_tag}.png"), dpi=600, bbox_inches="tight")
    plt.close(fig_pr)

    summary = pd.DataFrame(summary_rows).sort_values("roc_auc", ascending=False)
    print(f"\n=== {run_tag} ({len(feature_cols)} features, {int(y_true.sum())} positives) ===")
    print(summary.to_string(index=False))
    summary.to_csv(out_path(f"classic_ml_summary_{run_tag}.csv"), index=False)

    return summary


def plot_per_model_feature_comparison(target):
    """One ROC plot per model, overlaying the three feature sets (noloop/loopflag/withloop).
    """
    colors = {"noloop": "#9a9a96", "loopflag": "#2a78d6", "withloop": "#e34948"}
    labels = {"noloop": "No hairpin features", "loopflag": "Hairpin flag (TC at 3' loop end)",
              "withloop": "Hairpin properties: stem & loop lengths, energy"}

    for model_name in make_models(matched[target].values).keys():
        fig, ax = plt.subplots(figsize=(5.5, 5))
        lines = []  # (line, auc) -- legend sorted by AUC descending below
        for tag in FEATURE_SETS:
            pred = pd.read_csv(out_path(f"predictions_{model_name}_{tag}_{target}.csv"))
            y_true = pred[target].values
            prob = pred["pred_prob"].values
            fpr, tpr, _ = roc_curve(y_true, prob)
            auc = roc_auc_score(y_true, prob)
            (line,) = ax.plot(fpr, tpr, color=colors[tag], label=f"{labels[tag]} (AUC={auc:.3f})")
            lines.append((line, auc))

        ax.plot([0, 1], [0, 1], "--", color="lightgrey", zorder=0)
        ax.set_xlabel("False positive rate"); ax.set_ylabel("True positive rate")
        #ax.set_title(f"{model_name} — matched pairs, {target}\nROC by feature set (LOO-CV OOF)")
        lines.sort(key=lambda t: t[1], reverse=True)
        ax.legend([l for l, _ in lines], [l.get_label() for l, _ in lines], loc="lower right", fontsize=8)
        fig.tight_layout()
        fig.savefig(out_path(f"roc_by_featureset_{model_name}_{target}.png"), dpi=600, bbox_inches="tight")
        plt.close(fig)
        print(f"saved roc_by_featureset_{model_name}_{target}.png")


if __name__ == "__main__":
    summaries = [run_for_feature_set(tag, cols, tgt)
                 for tgt in TARGETS
                 for tag, cols in FEATURE_SETS.items()]
    combined = pd.concat(summaries, ignore_index=True)
    combined.to_csv(out_path("classic_ml_summary_matched_pairs_combined.csv"), index=False)
    print("\n=== combined: all targets x feature sets ===")
    print(combined.sort_values(["target", "feature_set", "roc_auc"], ascending=[True, True, False]).to_string(index=False))

    for tgt in TARGETS:
        plot_per_model_feature_comparison(tgt)
