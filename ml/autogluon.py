import pandas as pd
import mlflow
from autogluon.tabular import TabularPredictor
from sklearn.metrics import (precision_recall_curve, roc_curve, roc_auc_score, average_precision_score,
    confusion_matrix, precision_score,)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl.styles import PatternFill

df = pd.read_csv("/Users/mar/BIO/PROJECTS/MPOX/COMMENT/APOBEC3-Mpox-Hairpins/ml/ml_table_all.tsv", sep='\t')
df = pd.get_dummies(df, columns=['third_nt'], dtype="uint8")
#df = pd.get_dummies(df, columns=['minus1_nt'], dtype="uint8")
df = df[((df.loop_boundary_length == 3) & (df.loop_boundary_stem_length == 7)) |
        ((df.loop_boundary_length == 3) & (df.loop_boundary_stem_length == 6)) |
        ((df.loop_boundary_length == 3) & (df.loop_boundary_stem_length == 5)) |
        ((df.loop_boundary_length == 4) & (df.loop_boundary_stem_length == 6)) |
        ((df.loop_boundary_length == 4) & (df.loop_boundary_stem_length == 5))]

feature_set1 = ["grantham",
                "third_nt_A","third_nt_C","third_nt_G","third_nt_T",#"third_nt",
                #"minus1_nt_A","minus1_nt_C","minus1_nt_G","minus1_nt_T",#"is_YTCA","is_RTCA","is_TCW",
                #"in_stem","stem_length","stem_loop_length","stem_pin_energy","stem_linear_energy","stem_nn_energy","stem_cruciform_hairpin_dG","stem_linear_duplex_dG","stem_cruciform_relative_dG",
               "in_loop_boundary","loop_boundary_length","loop_boundary_stem_length","loop_boundary_pin_energy","loop_boundary_nn_energy"#,"loop_boundary_unpaired_prob"#,"loop_boundary_linear_energy"#,"loop_boundary_cruciform_hairpin_dG","loop_boundary_linear_duplex_dG","loop_boundary_cruciform_relative_dG"#,
               #"in_loop_other","loop_other_length","loop_other_stem_length","loop_other_pin_energy","loop_other_linear_energy","loop_other_nn_energy","loop_other_cruciform_hairpin_dG","loop_other_linear_duplex_dG","loop_other_cruciform_relative_dG"
               ]

print(df[feature_set1].dtypes)

target = "ds2_mutation"
#preset = "interpretable"
preset = "medium"
metric = "roc_auc"

def log_autogluon(predictor, test_df=None):
    lb = predictor.leaderboard()
    for _, r in lb.iterrows():
        mlflow.log_metric(f"score_val__{r['model'].replace('.', '_')}", r["score_val"])
        #mlflow.log_metric(f"roc_auc_val__{r['model'].replace('.', '_')}", r["score_val_roc_auc"])
    mlflow.log_metric(f"best_score_val", lb.iloc[0]["score_val"])
    #mlflow.log_metric(f"best_rocauc_val", lb.iloc[0]["score_val_roc_auc"])
    mlflow.log_param("best_model", lb.iloc[0]["model"])
    lb.to_csv("leaderboard.csv", index=False)
    mlflow.log_artifact("leaderboard.csv")

def save_predictions_xlsx(out, target, path="predictions.xlsx"):
    # добавим явную колонку типа ошибки — удобно фильтровать в Excel
    def err_type(r):
        if r["pred_class"] == 1 and r[target] == 0: return "FP"
        if r["pred_class"] == 0 and r[target] == 1: return "FN"
        return ""
    out = out.copy()
    out["error"] = out.apply(err_type, axis=1)

    out.to_excel(path, index=False)

    # подсветка через openpyxl
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active

    fp_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # красный
    fn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # жёлтый

    err_col = out.columns.get_loc("error") + 1   # openpyxl считает с 1
    for row in range(2, ws.max_row + 1):          # строка 1 — заголовок
        val = ws.cell(row=row, column=err_col).value
        fill = fp_fill if val == "FP" else fn_fill if val == "FN" else None
        if fill:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill

    wb.save(path)
    return path

"""
SHAP-интерпретация модели AutoGluon с логированием в MLflow.

Функция shap_explain_autogluon():
  - сама выбирает лучшую tree-модель из leaderboard;
  - считает SHAP (быстрый TreeExplainer; при неудаче — model-agnostic KernelExplainer);
  - строит summary plot и dependence-графики по топ-признакам;
  - логирует всё артефактами в MLflow.

Зависимости: pip install shap matplotlib mlflow
"""

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import shap
import mlflow


# tree-модели AutoGluon, для которых работает быстрый TreeExplainer
_TREE_KEYS = ("LightGBM", "XGBoost", "CatBoost", "RandomForest", "ExtraTrees", "RF", "XT")


def _pick_best_tree_model(predictor):
    """Возвращает имя лучшей (по score_val) tree-модели из leaderboard."""
    lb = predictor.leaderboard(silent=True) if "silent" in \
        predictor.leaderboard.__code__.co_varnames else predictor.leaderboard()
    lb = lb.sort_values("score_val", ascending=False)
    for name in lb["model"]:
        if any(k.lower() in name.lower() for k in _TREE_KEYS):
            return name
    return None


def _get_shap_values_tree(booster, X_trans):
    """SHAP для tree-модели; нормализует вывод к массиву (n_samples, n_features)."""
    explainer = shap.TreeExplainer(booster)
    sv = explainer.shap_values(X_trans)
    # для бинарной классификации разные версии shap отдают list из 2 массивов
    if isinstance(sv, list):
        sv = sv[1] if len(sv) == 2 else sv[0]
    # новые версии shap могут вернуть 3D (n, features, classes)
    if isinstance(sv, np.ndarray) and sv.ndim == 3:
        sv = sv[:, :, -1]
    return sv


def _get_shap_values_kernel(predictor, features, X_explain, background_n=None):
    """Model-agnostic fallback через predict_proba (медленно, но всегда работает)."""
    n = len(X_explain) if background_n is None else min(background_n, len(X_explain))
    background = shap.sample(X_explain, n, random_state=0)

    def f(X):
        return predictor.predict_proba(
            pd.DataFrame(X, columns=features)).iloc[:, 1].values

    explainer = shap.KernelExplainer(f, background)
    return explainer.shap_values(X_explain)


def shap_explain_autogluon(
    predictor,
    df,
    model_name=None,
    max_explain=200,
    top_k_dependence=4,
    log_to_mlflow=True,
    use_kernel=False
):
    """
    predictor       — обученный TabularPredictor
    df              — данные для объяснения (те же признаки, что при обучении)
    model_name      — имя модели из leaderboard; если None — берётся лучшая tree-модель
    max_explain     — сколько строк объяснять (SHAP на всех может быть медленным)
    top_k_dependence— сколько dependence-графиков построить (по важности)
    log_to_mlflow   — логировать артефакты в активный MLflow run
    Возвращает: (shap_values, X_used, feature_names)
    """
    features = predictor.feature_metadata_in.get_features()
    X_raw = df[features]
    if len(X_raw) > max_explain:
        X_raw = X_raw.sample(max_explain, random_state=0)

    # --- выбор модели ---
    if model_name is None:
        model_name = _pick_best_tree_model(predictor)

    used_method = None
    shap_values = None
    X_used = None

    # --- попытка 1: быстрый TreeExplainer на трансформированных признаках ---
    if model_name is not None and not use_kernel:
        try:
            booster = predictor._trainer.load_model(model_name).model
            X_trans = predictor.transform_features(X_raw, model=model_name) \
                if "model" in predictor.transform_features.__code__.co_varnames \
                else predictor.transform_features(X_raw)
            shap_values = _get_shap_values_tree(booster, X_trans)
            X_used = X_trans
            used_method = f"TreeExplainer[{model_name}]"
        except Exception as e:
            print(f"TreeExplainer не сработал ({e}); перехожу на KernelExplainer.")

    # --- попытка 2: model-agnostic fallback ---
    if shap_values is None:
        shap_values = _get_shap_values_kernel(predictor, features, X_raw)
        X_used = X_raw
        used_method = "KernelExplainer[predictor]"

    feature_names = list(X_used.columns)
    print(f"SHAP посчитан методом: {used_method}, объяснено строк: {len(X_used)}")

    # --- summary plot ---
    plt.figure()
    shap.summary_plot(shap_values, X_used, feature_names=feature_names, show=False)
    plt.title(f"SHAP summary — {used_method}")
    plt.tight_layout()
    plt.savefig("shap_summary.png", dpi=150, bbox_inches="tight")
    plt.close()

    # --- bar plot (глобальная важность) ---
    plt.figure()
    shap.summary_plot(shap_values, X_used, feature_names=feature_names,
                      plot_type="bar", show=False)
    plt.tight_layout()
    plt.savefig("shap_importance_bar.png", dpi=150, bbox_inches="tight")
    plt.close()

    # --- dependence plots по топ-признакам ---
    mean_abs = np.abs(shap_values).mean(axis=0)
    top_idx = np.argsort(mean_abs)[::-1][:top_k_dependence]
    dep_files = []
    for i in top_idx:
        fname = f"shap_dependence_{feature_names[i]}.png"
        plt.figure()
        shap.dependence_plot(int(i), shap_values, X_used,
                             feature_names=feature_names, show=False)
        plt.tight_layout()
        plt.savefig(fname, dpi=150, bbox_inches="tight")
        plt.close()
        dep_files.append(fname)

    # --- таблица важностей ---
    imp = pd.DataFrame({"feature": feature_names, "mean_abs_shap": mean_abs}) \
        .sort_values("mean_abs_shap", ascending=False)
    imp.to_csv("shap_importance.csv", index=False)

    # --- логирование в MLflow ---
    if log_to_mlflow:
        for fpath in ["shap_summary.png", "shap_importance_bar.png",
                      "shap_importance.csv"] + dep_files:
            mlflow.log_artifact(fpath)
        mlflow.log_param("shap_method", used_method)
        mlflow.log_param("shap_n_explained", len(X_used))

    return shap_values, X_used, feature_names


# ------------------------------------------------------------------
# Пример использования:
#
# from autogluon.tabular import TabularPredictor
# predictor = TabularPredictor(label="any_mutation", eval_metric="roc_auc",
#                              sample_weight="balance_weight").fit(
#     df[feature_set1 + ["any_mutation"]], presets="high")
#
# import mlflow
# mlflow.set_experiment("mpox-apobec3-hairpins")
# with mlflow.start_run(run_name="shap"):
#     shap_explain_autogluon(predictor, df)
# ------------------------------------------------------------------






mlflow.set_experiment("autogluon-ds2-newenergy")

with mlflow.start_run():
    mlflow.log_params({
        "target": target,
        "eval_metric": metric,
        "presets": preset,
        "n_features": len(feature_set1),
        "n_rows": len(df),
    })
    mlflow.log_param("features", ",".join(feature_set1))

    predictor = TabularPredictor(label=target, eval_metric=metric, problem_type='binary', positive_class=1).fit(df[feature_set1 + [target]], presets=preset)
    #predictor = TabularPredictor(label=target).fit(df[feature_set1 + [target]], presets=preset)
   
    
    ## Print interpretable models
    #model = predictor._trainer.load_model("RuleFit")
    #est = model.model
    #rules = est._get_rules()
    #print(rules.to_string())   
    #print("intercept:", getattr(est, "intercept_", None))
    #print("coef:", getattr(est, "coef_", None))
    
    #tree = predictor._trainer.load_model("Figs").model
    #print(tree)

    log_autogluon(predictor)

    y_true = df[target].values
    prob = predictor.predict_proba(df).iloc[:, -1].values

    threshold = prob[y_true == 1].min()
    y_pred = (prob >= threshold).astype(int)

    tn, fp, fn, tp = confusion_matrix(y_true, y_pred).ravel()
    recall = tp / (tp + fn)
    precision = precision_score(y_true, y_pred, zero_division=0)
    roc_auc = roc_auc_score(y_true, prob)
    pr_auc = average_precision_score(y_true, prob)

    mlflow.log_params({"threshold": float(threshold)})
    mlflow.log_metrics({
        "recall": recall, "precision": precision,
        "roc_auc": roc_auc, "pr_auc": pr_auc,
        "TP": tp, "FP": fp, "FN": fn, "TN": tn,
        "n_positives": int(tp + fn),
    })
                       
    odf = df.copy() 
    odf["pred_class"] = y_pred
    odf["pred_prob"] = prob
 
    run = mlflow.active_run()
    run_name = run.data.tags.get("mlflow.runName", run.info.run_id[:8])

    odf.to_csv(f"predictions_{run_name}.csv", index=False)
    mlflow.log_artifact(f"predictions_{run_name}.csv")
    save_predictions_xlsx(odf, target, f"predictions_{run_name}.xlsx")

    prec, rec, _ = precision_recall_curve(y_true, prob)
    fig, ax = plt.subplots(figsize=(6, 5))
    ax.plot(rec, prec, label=f"PR (AP = {pr_auc:.3f})")
    ax.scatter([recall], [precision], color="red", zorder=5,
               label=f"порог={threshold:.3f}\nrecall={recall:.2f}, prec={precision:.2f}")
    ax.set_xlabel("Recall"); ax.set_ylabel("Precision")
    ax.set_title("PR-кривая (рабочая точка = recall 100%)")
    ax.legend(loc="lower left")
    mlflow.log_figure(fig, "pr_curve.png")
    plt.close(fig)

    fpr, tpr, _ = roc_curve(y_true, prob)
    fig2, ax2 = plt.subplots(figsize=(6, 5))
    ax2.plot(fpr, tpr, label=f"ROC (AUC = {roc_auc:.3f})")
    ax2.plot([0, 1], [0, 1], "--", color="grey")
    ax2.set_xlabel("FPR"); ax2.set_ylabel("TPR"); ax2.set_title("ROC")
    ax2.legend(loc="lower right")
    mlflow.log_figure(fig2, "roc_curve.png")
    plt.close(fig2)

    shap_explain_autogluon(predictor, df, use_kernel=True)

